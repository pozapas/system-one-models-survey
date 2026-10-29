"""Build the second-rating workbooks for the version 1.2 sweep: every candidate, every graded study, at least one
extracted value per study, and a share of the census head codes, split over N independent raters.

Each rater gets one workbook with four sheets:
  Inclusion    a share of the 249 arXiv candidates of candidates_v12.csv, judged include or exclude;
  RiskOfBias   a share of the 215 studies graded in parts/rob_G*.csv, rated on the five domains;
  Extraction   ledger values from that rater's own risk-of-bias studies (one per study, then more at random),
               read from the cited source;
  Heads        a share of a sample of the census originals, coded by decision-head type.
Every candidate and every study goes to exactly one rater. The first-rater decisions go only to the key file,
which the raters never see. second_rater_agreement.py computes agreement from the returned workbooks.

Usage (needs openpyxl):  python make_second_rating_v12.py [--raters 10] [--rows 300] [--heads 38] [--seed 20261029]
"""
import argparse
import csv
import json
import random
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from make_second_rater import ANCHORS, DOMAINS

SURVEY = Path(__file__).resolve().parents[1]
PAPER4 = SURVEY.parent
LEDGER = PAPER4 / "shared" / "ledger"
CENSUS = PAPER4 / "shared" / "census"
OUT = SURVEY / "notes" / "second_rating_v12"
KEY = SURVEY / "notes" / "second_rating_v12_key"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HEAD_PANELS = ["a fixed slots", "b option markers", "c label logits", "d pair scoring", "e constrained decoding",
               "undisclosed"]
MIN_PER = {"inclusion": 1.0, "rob": 10.0, "extraction": 1.5, "heads": 2.5}   # minutes per item, for the guide


def sheet(wb, title, cols, rows, fill_cols, widths, validation=None):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, col in enumerate(cols, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths[i - 1]
        if col in fill_cols:
            for j in range(2, len(rows) + 2):
                ws.cell(row=j, column=i).fill = YELLOW
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if validation and rows:
        col, options = validation
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        letter = ws.cell(row=1, column=cols.index(col) + 1).column_letter
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "A2"
    return ws


def split(items, n):
    return [items[i::n] for i in range(n)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raters", type=int, default=10)
    ap.add_argument("--rows", type=int, default=300)
    ap.add_argument("--heads", type=int, default=38)
    ap.add_argument("--seed", type=int, default=20261029)
    a = ap.parse_args()
    rng = random.Random(a.seed)

    cand = list(csv.DictReader((LEDGER / "candidates_v12.csv").open(encoding="utf-8")))
    rng.shuffle(cand)
    graded = set()
    for p in sorted((LEDGER / "parts").glob("rob_G*.csv")):
        graded |= {r["study_key"] for r in csv.DictReader(p.open(encoding="utf-8-sig"))}
    rob = [r for r in csv.DictReader((LEDGER / "rob.csv").open(encoding="utf-8")) if r["study_key"] in graded]
    rng.shuffle(rob)
    led = [r for r in csv.DictReader((LEDGER / "ledger.csv").open(encoding="utf-8"))
           if r["ledger_version"].strip() == "1.2" and r["value"].strip() and r["location"].strip()]
    by_study = defaultdict(list)
    for r in led:
        by_study[r["study_key"]].append(r)
    url = {}
    for r in csv.DictReader((LEDGER / "ledger.csv").open(encoding="utf-8")):
        url.setdefault(r["study_key"], r["source_url"])
    heads = list(csv.DictReader((CENSUS / "census_heads.csv").open(encoding="utf-8")))
    by_panel = defaultdict(list)
    for h in heads:
        by_panel[h["panel"]].append(h)
    head_sample = []
    for p, items in sorted(by_panel.items()):   # proportional to each head type, at least one per non-empty type
        k = max(1, round(a.heads * len(items) / len(heads)))
        head_sample += rng.sample(items, min(k, len(items)))
    rng.shuffle(head_sample)

    cand_s, rob_s, head_s = split(cand, a.raters), split(rob, a.raters), split(head_sample, a.raters)
    # extraction: one row per study, then extra rows at random from the same rater's studies up to the target
    extra_total = max(0, a.rows - len(rob))
    key = {"seed": a.seed, "raters": {}}
    OUT.mkdir(parents=True, exist_ok=True)
    KEY.mkdir(parents=True, exist_ok=True)
    summary = []
    for i in range(a.raters):
        rid = f"R{i + 1:02d}"
        studies = rob_s[i]
        ext = []
        for s in studies:
            rows = by_study.get(s["study_key"], [])
            if rows:
                ext.append(rng.choice(rows))
        pool = [r for s in studies for r in by_study.get(s["study_key"], []) if r not in ext]
        ext += rng.sample(pool, min(len(pool), round(extra_total / a.raters)))
        ext.sort(key=lambda r: (r["study_key"], r["row_id"]))

        wb = Workbook()
        ins = wb.active
        ins.title = "Instructions"
        n = {"inclusion": len(cand_s[i]), "rob": len(studies), "extraction": len(ext), "heads": len(head_s[i])}
        minutes = sum(n[k] * MIN_PER[k] for k in n)
        lines = [
            f"Second rating of the survey 'System One Decision Models', workbook {rid}.",
            f"Your share: {n['inclusion']} screening decisions, {n['rob']} studies to grade on five domains, "
            f"{n['extraction']} values to read from those same studies, and {n['heads']} decision-head codes. "
            f"Expected time: about {minutes / 60:.1f} hours.",
            "Read the guide (Second_rating_guide.docx) before you start. Work alone, do not discuss items with the "
            "other raters, and do not open the survey, its tables or its repository.",
            "Fill only the yellow cells. Leave a cell empty if you cannot decide and explain why in the notes column.",
            "",
            "Rating anchors for the RiskOfBias sheet:",
        ] + [f"{name}: {ANCHORS[k]}" for k, name, _ in DOMAINS] + [
            "Domain 5 also: a study that mixes new items with public sets released before 2026 is unclear.",
            "",
            "Decision-head codes for the Heads sheet: a fixed slots (a fixed number of output slots); b option markers "
            "(each option gets a marker or pointer position that a head scores, or a small head scores all options "
            "from the final hidden states); c label logits (option letters or label tokens read from the language-model "
            "head, or a Yes-minus-No readout); d pair scoring (each state-question-option scored separately, as a "
            "cross-encoder or NLI model); e constrained decoding (the answer generated under a grammar or option mask); "
            "undisclosed (the card and code do not say).",
            "",
            "Save the workbook under its own name and return it when done.",
        ]
        for line in lines:
            ins.append([line])
        ins.column_dimensions["A"].width = 150
        for row in ins.iter_rows():
            row[0].alignment = Alignment(wrap_text=True)
        ins["A1"].font = Font(bold=True, size=13)

        sheet(wb, "Inclusion", ["arxiv_id", "title", "abstract", "link", "decision", "notes"],
              [[c["arxiv_id"], c["title"], c["abstract"], f"https://arxiv.org/abs/{c['arxiv_id']}", "", ""]
               for c in cand_s[i]],
              {"decision", "notes"}, [14, 40, 90, 32, 12, 30], ("decision", ["include", "exclude"]))
        rob_rows = []
        for s in studies:
            for k, name, _ in DOMAINS:
                rob_rows.append([s["study_key"], url.get(s["study_key"], ""), name, "", ""])
        sheet(wb, "RiskOfBias", ["study", "source", "domain", "rating", "justification"], rob_rows,
              {"rating", "justification"}, [26, 40, 28, 12, 70], ("rating", ["low", "unclear", "high"]))
        sheet(wb, "Extraction",
              ["item", "study", "source", "location", "model", "dataset", "metric", "detail", "value", "notes"],
              [[j + 1, r["study_key"], r["source_url"], r["location"], r["model"], r["dataset"], r["metric"],
                r["value_key"].split(".", 1)[-1], "", ""] for j, r in enumerate(ext)],
              {"value", "notes"}, [6, 24, 36, 26, 28, 30, 28, 30, 12, 30])
        sheet(wb, "Heads", ["item", "repository", "card", "head", "evidence", "notes"],
              [[j + 1, h["id"], f"https://huggingface.co/{h['id']}", "", "", ""] for j, h in enumerate(head_s[i])],
              {"head", "evidence", "notes"}, [6, 44, 52, 22, 60, 30], ("head", HEAD_PANELS))
        wb.save(OUT / f"second_rating_{rid}.xlsx")

        key["raters"][rid] = {
            "inclusion": {c["arxiv_id"]: c["screen_decision"] for c in cand_s[i]},
            "rob": {s["study_key"]: {name: s[f"{k}_rating"] for k, name, _ in DOMAINS} for s in studies},
            "extraction": {str(j + 1): {"row_id": r["row_id"], "value": r["value"]} for j, r in enumerate(ext)},
            "heads": {str(j + 1): {"id": h["id"], "panel": h["panel"]} for j, h in enumerate(head_s[i])},
        }
        summary.append((rid, n, minutes))
    (KEY / "first_rating_key.json").write_text(json.dumps(key, indent=1), encoding="utf-8")
    tot = defaultdict(int)
    for rid, n, minutes in summary:
        print(f"{rid}: {n}  ~{minutes / 60:.1f} h")
        for k, v in n.items():
            tot[k] += v
    print("totals:", dict(tot), "| key:", KEY / "first_rating_key.json")


if __name__ == "__main__":
    main()
