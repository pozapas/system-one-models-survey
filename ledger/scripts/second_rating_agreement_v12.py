"""Agreement between the first rating and the returned second-rating workbooks of the version 1.2 sweep.

Reads the ten workbooks made by make_second_rating_v12.py from a folder, checks that only the yellow cells changed,
and compares each answer with notes/second_rating_v12_key/first_rating_key.json. Screening, risk of bias and head
codes get raw agreement and Cohen's kappa (risk of bias also a linearly weighted kappa on low < unclear < high, overall
and per domain). Extraction uses same_value() of second_rater_agreement.py. Writes a json of values and a csv of every
disagreement for adjudication.

Usage (needs openpyxl):  python second_rating_agreement_v12.py <returned folder> [out.json] [disagreements.csv]
"""
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from second_rater_agreement import ORDER, kappa, rows, same_value

SURVEY = Path(__file__).resolve().parents[1]
KEY = SURVEY / "notes" / "second_rating_v12_key" / "first_rating_key.json"
ORIG = SURVEY / "notes" / "second_rating_v12"
FILL = {"Inclusion": {"decision", "notes"}, "RiskOfBias": {"rating", "justification"},
        "Extraction": {"value", "notes"}, "Heads": {"head", "evidence", "notes"}}


def fixed_cells_changed(orig, back):
    changed = []
    for name, fill in FILL.items():
        a, b = orig[name], back[name]
        head = [c.value for c in a[1]]
        for ra, rb in zip(a.iter_rows(min_row=2, values_only=True), b.iter_rows(min_row=2, values_only=True)):
            for col, x, y in zip(head, ra, rb):
                if col not in fill and (x or "") != (y or ""):
                    changed.append((name, col, x, y))
        if a.max_row != b.max_row:
            changed.append((name, "rows", a.max_row, b.max_row))
    return changed


def stats(pairs, weights=None):
    if not pairs:
        return {"n": 0}
    a, b = zip(*pairs)
    k = kappa(a, b)
    out = {"n": len(pairs), "agree": sum(x == y for x, y in pairs),
           "agreement": round(sum(x == y for x, y in pairs) / len(pairs), 3),
           "kappa": None if k is None else round(k, 2)}
    if weights:
        out["weightedKappa"] = round(kappa(a, b, weights=weights), 2)
    return out


def main():
    # optional --key <file> --orig <folder> score another round (for example make_second_rating_v11rest.py)
    args = sys.argv[1:]
    key_path, orig = KEY, ORIG
    for flag in ("--key", "--orig"):
        if flag in args:
            i = args.index(flag)
            if flag == "--key":
                key_path = Path(args[i + 1])
            else:
                orig = Path(args[i + 1])
            del args[i:i + 2]
    folder = Path(args[0])
    out = Path(args[1]) if len(args) > 1 else folder / "second_rating_v12.json"
    dis_path = Path(args[2]) if len(args) > 2 else folder / "disagreements_v12.csv"
    key = json.loads(key_path.read_text(encoding="utf-8"))["raters"]
    inc, rob, ext, heads, dis, blanks, integrity = [], [], [], [], [], defaultdict(int), {}
    rob_dom = defaultdict(list)
    for rid, k in key.items():
        back = load_workbook(folder / f"second_rating_{rid}.xlsx")
        integrity[rid] = fixed_cells_changed(load_workbook(orig / f"second_rating_{rid}.xlsx"), back)
        for r in rows(back["Inclusion"]):
            d = str(r.get("decision") or "").strip().lower()
            if not d:
                blanks["inclusion"] += 1
                continue
            first = k["inclusion"][r["arxiv_id"]]
            inc.append((first, d))
            if first != d:
                dis.append([rid, "inclusion", r["arxiv_id"], "", first, d, r.get("notes") or ""])
        for r in rows(back["RiskOfBias"]):
            d = str(r.get("rating") or "").strip().lower()
            if not d:
                blanks["rob"] += 1
                continue
            first = k["rob"][r["study"]][r["domain"]]
            first = "unclear" if first == "some_concerns" else first
            rob.append((first, d))
            rob_dom[r["domain"]].append((first, d))
            if first != d:
                dis.append([rid, "rob", r["study"], r["domain"], first, d, r.get("justification") or ""])
        for r in rows(back["Extraction"]):
            v = r.get("value")
            if v in (None, ""):
                blanks["extraction"] += 1
                continue
            first = k["extraction"][str(r["item"])]
            ok = same_value(first["value"], v)
            ext.append(ok)
            if not ok:
                dis.append([rid, "extraction", f"{r['study']} item {r['item']}", first["row_id"], first["value"], v,
                            r.get("notes") or ""])
        for r in rows(back["Heads"]):
            d = str(r.get("head") or "").strip()
            if not d:
                blanks["heads"] += 1
                continue
            d = d.split()[0] if d != "undisclosed" else d
            first = k["heads"][str(r["item"])]
            heads.append((first["panel"], d))
            if first["panel"] != d:
                dis.append([rid, "heads", first["id"], "", first["panel"], d, r.get("evidence") or ""])
    res = {"inclusion": stats(inc), "rob": stats(rob, ORDER),
           "robByDomain": {dname: stats(p, ORDER) for dname, p in rob_dom.items()},
           "extraction": {"n": len(ext), "agree": sum(ext), "agreement": round(sum(ext) / len(ext), 3) if ext else None},
           "heads": stats(heads), "blank": dict(blanks),
           "fixedCellsChanged": {r: len(c) for r, c in integrity.items()}}
    out.write_text(json.dumps(res, indent=1), encoding="utf-8")
    with dis_path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["rater", "task", "item", "detail", "first", "second", "second_note"])
        w.writerows(dis)
    print(json.dumps(res, indent=1))
    for r, c in integrity.items():
        for x in c[:5]:
            print("changed", r, x)
    print(len(dis), "disagreements ->", dis_path)


if __name__ == "__main__":
    main()
